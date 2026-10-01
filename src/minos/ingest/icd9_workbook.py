"""Separately attributed NFZ workbook hierarchy; never joined into CZS facts."""

import heapq
import re
from collections import defaultdict
from collections.abc import Iterator
from dataclasses import dataclass, field
from typing import TYPE_CHECKING, Annotated, BinaryIO, ClassVar, cast
from urllib.parse import urlsplit

from pydantic import BaseModel, ConfigDict, Field, ValidationError

from minos.domain.normalization import normalize_description
from minos.domain.reference import (
    ConceptDescription,
    ReferenceNode,
    ReferenceRelation,
    SourceError,
    SourceHeader,
)
from minos.ingest.icd10 import ReferenceEvent
from minos.ingest.reference_manifest import SourceManifest
from minos.ingest.workbook_archive import (
    DEFAULT_WORKBOOK_LIMITS,
    WorkbookLimits,
    WorksheetLayout,
    preflight_workbook,
)

if TYPE_CHECKING:
    from openpyxl.cell.read_only import EmptyCell, ReadOnlyCell
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet
    from openpyxl.worksheet.worksheet import Worksheet

_CODE = re.compile(r"(?:[A-Z]{1,2}|[A-Z][0-9]{2,3}|[0-9]{2,3}(?:\.[0-9]{1,5})?)\Z")
_MAIN_HEADERS = (
    "Rozdział",
    None,
    "Podrozdział",
    None,
    "Kategoria główna",
    None,
    "Kategoria szczegółowa",
    None,
)
_SUBHEADERS = ("nr ", "tytuł", "nr", "tytuł", "nr", "tytuł", "nr", "tytuł")
_NOTES = {
    "wł_wył_rozdz": ("Nr:", "TD:", "Wł/Wył:", "tytuł"),
    "wł_wył. podrozdz": ("nr", "Td:", "wł/wył", "tytuł"),
    "wł_wył_kat gl": ("Numer:", "TD:", "Wł/Wył:", "tytuł"),
    "Uwagi": ("Kod procedury", "Tytuł", "Techniki, uwagi", None),
}


class _Pair(BaseModel):
    model_config: ClassVar[ConfigDict] = ConfigDict(
        extra="forbid", strict=True, frozen=True
    )
    code: str | int
    name: Annotated[str, Field(min_length=1, max_length=65_536)]


@dataclass
class _Entry:
    code: str
    source_code: str
    name: str
    level: int
    cell: str
    cell_type: str
    rank: int
    variants: set[tuple[str, str, str]] = field(default_factory=set)


@dataclass(frozen=True)
class _Evidence:
    code: str
    text: str
    attributes: tuple[tuple[str, str], ...]


def _selection(manifest: SourceManifest) -> str:
    try:
        url = urlsplit(manifest.source_url)
    except ValueError:
        raise SourceError("workbook manifest URL is invalid") from None
    prefix = "/download/gfx/nfz/pl/defaultstronaopisowa/782/1/2/"
    if manifest.system != "icd9-pl" or manifest.origin != "nfz-workbook":
        raise SourceError("workbook manifest adapter conflicts")
    if (
        url.scheme != "https"
        or url.netloc != "www.nfz.gov.pl"
        or url.query
        or url.fragment
        or not url.path.startswith(prefix)
    ):
        raise SourceError("workbook manifest URL is unsupported")
    filename = url.path[len(prefix) :]
    if re.fullmatch(r"[a-zA-Z0-9_.-]{3,140}\.xlsx", filename) is None:
        raise SourceError("workbook manifest filename is unsupported")
    version = re.search(r"(?:_|-)([1-9][0-9]{0,2}\.[0-9]{1,3})_?\.xlsx\Z", filename)
    if version is None or version[1] != manifest.source_version:
        raise SourceError("workbook version conflicts with manifest")
    file_id, separator, selected = manifest.source_file_id.partition("#")
    if (
        separator != "#"
        or file_id != filename
        or selected not in ("ICD-9_PL ", "ICD-9_PL")
    ):
        raise SourceError("workbook requires an explicit reviewed worksheet identity")
    return selected


def _pair(code: object, name: object) -> tuple[str, _Pair]:
    try:
        value = _Pair.model_validate({"code": code, "name": name})
    except ValidationError:
        raise SourceError("workbook code/title cell type is unsupported") from None
    key = str(value.code).strip()
    if _CODE.fullmatch(key) is None or not value.name.strip():
        raise SourceError("workbook code/title cell value is unsupported")
    return key, value


def _rows(
    sheet: ReadOnlyWorksheet, layout: WorksheetLayout, limits: WorkbookLimits
) -> Iterator[tuple[ReadOnlyCell | EmptyCell, ...]]:
    # Explicit iteration limits also bound bogus declared dimensions and gaps.
    sheet.reset_dimensions()
    # The pinned reader reuses Worksheet.iter_rows; its stub constrains self.
    for row in cast("Worksheet", cast(object, sheet)).iter_rows(
        min_row=1,
        max_row=layout.rows,
        max_col=max(layout.columns, 8),
        values_only=False,
    ):
        for cell in row:
            if cell.data_type in ("f", "e", "d"):
                raise SourceError("workbook formula, error or date cell is unsupported")
            value = cell.value
            if isinstance(value, str) and len(value) > limits.max_text_characters:
                raise SourceError("workbook cell text exceeds limit")
        yield cast(tuple["ReadOnlyCell | EmptyCell", ...], row)


def _main_table(
    sheet: ReadOnlyWorksheet, layout: WorksheetLayout, limits: WorkbookLimits
) -> tuple[dict[str, _Entry], set[tuple[str, str]], list[_Evidence]]:
    entries: dict[str, _Entry] = {}
    edges: set[tuple[str, str]] = set()
    evidence: list[_Evidence] = []
    if (
        set(layout.merges) != {"A1:B1", "C1:D1", "E1:F1", "G1:H1"}
        or len(layout.merges) != 4
    ):
        raise SourceError("workbook main header merges are unsupported")
    for number, row in enumerate(_rows(sheet, layout, limits), 1):
        values = tuple(cell.value for cell in row)
        if number <= 2:
            expected = _MAIN_HEADERS if number == 1 else _SUBHEADERS
            if values[:8] != expected or any(value is not None for value in values[8:]):
                raise SourceError("workbook main header layout is unsupported")
            continue
        if any(value not in (None, "") for value in values[8:]):
            raise SourceError("workbook main table has unexpected populated columns")
        parent: str | None = None
        for level in range(4):
            code, name = values[2 * level : 2 * level + 2]
            if code in (None, "") and name in (None, ""):
                continue
            key, pair = _pair(code, name)
            cell = "ACEG"[level] + str(number)
            _record(entries, evidence, key, pair, level + 1, cell, sheet.title, limits)
            if parent is not None:
                edges.add((parent, key))
                if len(edges) > limits.max_nodes:
                    raise SourceError("workbook hierarchy edges exceed limit")
            parent = key
    if not entries:
        raise SourceError("workbook classification table is empty")
    return entries, edges, evidence


def _record(
    entries: dict[str, _Entry],
    evidence: list[_Evidence],
    key: str,
    pair: _Pair,
    level: int,
    cell: str,
    sheet: str,
    limits: WorkbookLimits,
) -> None:
    raw = str(pair.code)
    if key not in entries:
        if len(entries) + len(evidence) >= limits.max_nodes:
            raise SourceError("workbook node count exceeds limit")
        entries[key] = _Entry(
            key, raw, pair.name, level, cell, type(pair.code).__name__, len(entries)
        )
        entries[key].variants.add((raw, pair.name, type(pair.code).__name__))
    entry = entries[key]
    if entry.level != level:
        raise SourceError("workbook code has conflicting hierarchy levels")
    variant = (raw, pair.name, type(pair.code).__name__)
    if variant not in entry.variants:
        if len(entries) + len(evidence) >= limits.max_nodes:
            raise SourceError("workbook evidence count exceeds limit")
        evidence.append(
            _Evidence(
                key,
                pair.name,
                (
                    ("evidence_kind", "conflicting_title_or_code_form"),
                    ("source_code", raw),
                    ("worksheet", sheet),
                    ("cell", cell),
                    ("cell_type", type(pair.code).__name__),
                ),
            )
        )
        entry.variants.add(variant)


def _annotations(
    sheet: ReadOnlyWorksheet,
    layout: WorksheetLayout,
    entries: dict[str, _Entry],
    evidence: list[_Evidence],
    limits: WorkbookLimits,
) -> None:
    if layout.merges:
        raise SourceError("workbook annotation merges are unsupported")
    for number, row in enumerate(_rows(sheet, layout, limits), 1):
        values = tuple(cell.value for cell in row)
        if number == 1:
            if values[:4] != _NOTES[sheet.title] or any(
                value is not None for value in values[4:]
            ):
                raise SourceError("workbook annotation header layout is unsupported")
            continue
        if any(value not in (None, "") for value in values[4:]):
            raise SourceError("workbook annotation has unexpected populated columns")
        code, title, marker, text = values[:4]
        if all(value in (None, "") for value in (code, title, marker, text)):
            continue
        key, pair = _pair(code, title)
        if sheet.title == "Uwagi":
            text, marker = marker, "note"
        if text in (None, ""):
            continue
        if (
            not isinstance(text, str)
            or not isinstance(marker, str)
            or not marker.strip()
        ):
            raise SourceError("workbook annotation cell type is unsupported")
        if len(entries) + len(evidence) >= limits.max_nodes:
            raise SourceError("workbook annotation count exceeds limit")
        evidence.append(
            _Evidence(
                key,
                text,
                (
                    ("evidence_kind", marker),
                    ("source_code", str(pair.code)),
                    ("source_title", pair.name),
                    ("worksheet", sheet.title),
                    ("row", str(number)),
                ),
            )
        )


def _ordered(entries: dict[str, _Entry], edges: set[tuple[str, str]]) -> list[str]:
    # Stable Kahn ordering preserves first-seen order where containment allows it.
    children: dict[str, list[str]] = defaultdict(list)
    degree = dict.fromkeys(entries, 0)
    for parent, child in edges:
        children[parent].append(child)
        degree[child] += 1
    ready = [(entry.rank, key) for key, entry in entries.items() if not degree[key]]
    heapq.heapify(ready)
    ordered: list[str] = []
    while ready:
        _, key = heapq.heappop(ready)
        ordered.append(key)
        for child in children[key]:
            degree[child] -= 1
            if not degree[child]:
                heapq.heappush(ready, (entries[child].rank, child))
    if len(ordered) != len(entries):
        raise SourceError("workbook containment has a cycle")
    return ordered


def _events(
    entries: dict[str, _Entry],
    edges: set[tuple[str, str]],
    evidence: list[_Evidence],
    selected: str,
) -> Iterator[ReferenceEvent]:
    ordered = _ordered(entries, edges)
    ordinals = {key: number for number, key in enumerate(ordered, 1)}
    for key in ordered:
        entry = entries[key]
        yield ReferenceNode(
            ordinals[key],
            entry.source_code,
            entry.name,
            (
                ("level", str(entry.level)),
                ("worksheet", selected),
                ("cell", entry.cell),
                ("cell_type", entry.cell_type),
            ),
            "classification",
            key,
            (
                ConceptDescription(
                    "pl", "preferred", entry.name, normalize_description(entry.name)
                ),
            ),
        )
    for parent, child in sorted(
        edges, key=lambda edge: (ordinals[edge[0]], ordinals[edge[1]])
    ):
        yield ReferenceRelation(ordinals[parent], ordinals[child], "contains")
    for number, item in enumerate(evidence, len(entries) + 1):
        yield ReferenceNode(number, "", item.text, item.attributes, "text", None)
        if item.code in ordinals:
            yield ReferenceRelation(ordinals[item.code], number, "contains")


def parse_icd9_workbook(
    source: BinaryIO,
    manifest: SourceManifest,
    *,
    limits: WorkbookLimits = DEFAULT_WORKBOOK_LIMITS,
) -> Iterator[ReferenceEvent]:
    """Validate every ZIP part before reader construction; consume before commit."""
    selected = _selection(manifest)
    layouts = preflight_workbook(source, limits=limits)
    try:
        import openpyxl
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError:
        raise SourceError(
            "workbook reader requires the pinned reference extra"
        ) from None
    if openpyxl.__version__ != "3.1.5":
        raise SourceError("workbook reader version is incompatible")
    workbook = None
    try:
        workbook = openpyxl.load_workbook(
            source, read_only=True, data_only=False, keep_links=False
        )
        if (
            not workbook.read_only
            or len(workbook.sheetnames) > 16
            or len(set(workbook.sheetnames)) != len(workbook.sheetnames)
        ):
            raise SourceError("workbook worksheet identities are invalid")
        if selected not in workbook.sheetnames:
            raise SourceError("workbook selected worksheet is missing")
        allowed = {"ICD-9_PL ", "ICD-9_PL", *_NOTES}
        if not set(workbook.sheetnames) <= allowed:
            raise SourceError("workbook worksheet layout is unsupported")
        entries, edges, evidence = _main_table(
            cast("ReadOnlyWorksheet", workbook[selected]), layouts[selected], limits
        )
        for name in _NOTES:
            if name in workbook.sheetnames:
                _annotations(
                    cast("ReadOnlyWorksheet", workbook[name]),
                    layouts[name],
                    entries,
                    evidence,
                    limits,
                )
        yield SourceHeader(
            (("worksheet", selected), ("status_evidence", "not supplied by workbook")),
            (),
        )
        yield from _events(entries, edges, evidence, selected)
    except (
        ValueError,
        KeyError,
        TypeError,
        IndexError,
        OSError,
        InvalidFileException,
    ) as error:
        raise SourceError("workbook reader rejected source structure") from error
    finally:
        if workbook is not None:
            workbook.close()
