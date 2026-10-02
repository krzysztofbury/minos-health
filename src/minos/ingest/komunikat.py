"""Bounded, explicit 2026 NFZ komunikat layouts; no algorithm execution."""

import re
from collections.abc import Iterator
from typing import TYPE_CHECKING, Annotated, BinaryIO, ClassVar, cast
from zipfile import ZipFile

from defusedxml.ElementTree import fromstring
from pydantic import BaseModel, ConfigDict, Field, ValidationError

from minos.domain.komunikat import (
    AlgorithmMetadata,
    DepartmentSpecialty,
    GrouperType,
    ProcedureRank,
    first_reporting_month,
)
from minos.domain.reference import SourceError
from minos.ingest.komunikat_manifest import KomunikatManifest
from minos.ingest.workbook_archive import (
    DEFAULT_WORKBOOK_LIMITS,
    WorkbookLimits,
    preflight_docx,
    preflight_workbook,
)

if TYPE_CHECKING:
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet
    from openpyxl.worksheet.worksheet import Worksheet

WORD = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
RANK_HEADERS = (
    "Kod procedury ICD-9",
    "Nazwa procedury ICD-9\n",
    "Ranga procedury ICD-9",
)
SPECIALTY_HEADERS = (
    "Kod specjalności komórki organizacyjnej",
    "Kod świadczenia",
    "Kody ICD",
)


class _RankRow(BaseModel):
    model_config: ClassVar[ConfigDict] = ConfigDict(extra="forbid", strict=True)
    code: Annotated[
        str, Field(pattern=r"^(?:[A-Z][0-9]{2,3}|[0-9]{2,3}(?:\.[0-9]{1,5})?)$")
    ]
    title: Annotated[str, Field(min_length=1, max_length=65_536)]
    rank: int = Field(ge=0, le=4)


class _SpecialtyRow(BaseModel):
    model_config: ClassVar[ConfigDict] = ConfigDict(extra="forbid", strict=True)
    code: Annotated[str, Field(pattern=r"^[0-9]{4}$")]
    service_codes: Annotated[str, Field(min_length=1, max_length=65_536)]
    icd_codes: Annotated[str, Field(min_length=1, max_length=65_536)]


def parse_algorithm(
    source: BinaryIO,
    manifest: KomunikatManifest,
    *,
    limits: WorkbookLimits = DEFAULT_WORKBOOK_LIMITS,
) -> AlgorithmMetadata:
    preflight_docx(source, limits=limits)
    with ZipFile(source) as archive:
        relationships = fromstring(archive.read("_rels/.rels"))
        main = [
            item.attrib.get("Target")
            for item in relationships
            if item.attrib.get("Type")
            == "http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument"
        ]
        if main != ["word/document.xml"]:
            raise SourceError("algorithm main document relationship is unsupported")
        document = fromstring(archive.read("word/document.xml"))
    if document.tag != WORD + "document" or len(document.findall(WORD + "body")) != 1:
        raise SourceError("algorithm document body is unsupported")
    if any(
        element.tag
        in {
            WORD + "ins",
            WORD + "del",
            WORD + "moveFrom",
            WORD + "moveTo",
            WORD + "altChunk",
        }
        for element in document.iter()
    ):
        raise SourceError("algorithm revision or external body is unsupported")
    paragraphs = [
        "".join(text.text or "" for text in paragraph.iter(WORD + "t"))
        for paragraph in document.iter(WORD + "p")
    ]
    if len(paragraphs) > 10_000:
        raise SourceError("algorithm paragraph count exceeds limit")
    if "Opis działania algorytmu JGP" not in paragraphs or "X.Y.ZZZZ" not in paragraphs:
        raise SourceError("algorithm versioning layout is unsupported")
    _versioning_layout(paragraphs)
    return _algorithm_version(paragraphs, manifest.algorithm_version)


def _versioning_layout(paragraphs: list[str]) -> None:
    normalized = {" ".join(text.replace("\u2013", "-").split()) for text in paragraphs}
    required = {
        "Y - oznacza okres sprawozdawczy, od którego w danym roku obowiązuje publikowana wersja algorytmu grupera",
        "ZZZZ - oznacza rok, dla którego opublikowany został opis działania algorytmu grupera.",
        "Wersja algorytmu grupera JGP obowiązuje z dokładnością do okresu sprawozdawczego (miesiąca). Oznacza to, że dla danego okresu sprawozdawczego obowiązuje tylko jedna wersja algorytmu.",
    }
    if not required <= normalized:
        raise SourceError("algorithm reporting-month semantics are unsupported")


def _algorithm_version(paragraphs: list[str], expected: str) -> AlgorithmMetadata:
    versions = [
        (
            number,
            text,
            re.fullmatch(
                r"wersja ([1-9][0-9]?\.(?:[1-9]|1[0-2])\.20[0-9]{2})",
                text.strip(),
                re.IGNORECASE,
            ),
        )
        for number, text in enumerate(paragraphs, 1)
        if re.match(r"wersja [0-9]", text.strip(), re.IGNORECASE)
    ]
    if len(versions) != 1:
        raise SourceError("algorithm version identity is missing or ambiguous")
    number, text, match = versions[0]
    if match is None:
        raise SourceError("algorithm version identity is missing or ambiguous")
    version = match[1]
    if version != expected:
        raise SourceError("algorithm version conflicts with manifest")
    return AlgorithmMetadata(version, first_reporting_month(version), text, number)


def _rows(
    sheet: ReadOnlyWorksheet, rows: int
) -> Iterator[tuple[int, tuple[object, ...]]]:
    sheet.reset_dimensions()
    for number, cells in enumerate(
        cast("Worksheet", cast(object, sheet)).iter_rows(
            max_row=rows, max_col=3, values_only=False
        ),
        1,
    ):
        if any(cell.data_type in ("f", "e", "d") for cell in cells):
            raise SourceError("komunikat formula, error or date cell is unsupported")
        yield number, tuple(cell.value for cell in cells)


def parse_companion(
    source: BinaryIO,
    manifest: KomunikatManifest,
    *,
    ranks: bool,
    limits: WorkbookLimits = DEFAULT_WORKBOOK_LIMITS,
) -> Iterator[ProcedureRank | DepartmentSpecialty]:
    try:
        import openpyxl
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError:
        raise SourceError(
            "workbook reader requires the pinned reference extra"
        ) from None
    if openpyxl.__version__ != "3.1.5":
        raise SourceError("workbook reader version is incompatible")

    layouts = preflight_workbook(source, limits=limits)
    expected = (
        cast(dict[GrouperType, int], manifest.expected_ranks)
        if ranks
        else manifest.expected_specialties
    )
    year = manifest.algorithm_version.split(".")[2]
    if set(layouts) != {scope + " " + year for scope in expected}:
        raise SourceError("komunikat worksheet identities conflict")
    book = None
    try:
        book = openpyxl.load_workbook(
            source, read_only=True, data_only=False, keep_links=False
        )
        if set(book.sheetnames) != set(layouts) or not book.read_only:
            raise SourceError("workbook reader worksheet identities conflict")
        for sheet in book:
            scope = cast(GrouperType, sheet.title.split(" ")[0])
            layout = layouts[sheet.title]
            if layout.columns != 3 or layout.merges or layout.rows < 4:
                raise SourceError("komunikat table layout is unsupported")
            count = 0
            seen: set[str] = set()
            for number, values in _rows(sheet, layout.rows):
                if number <= 3:
                    _header(values, number, scope, ranks=ranks)
                    continue
                if all(value is None for value in values):
                    continue
                event = _companion_row(values, scope, sheet.title, number, ranks=ranks)
                if event.code in seen:
                    raise SourceError("komunikat code is duplicated within a worksheet")
                seen.add(event.code)
                count += 1
                yield event
            if count != expected[scope]:
                raise SourceError("reviewed komunikat row count conflicts")
    except ValueError, KeyError, TypeError, IndexError, OSError, InvalidFileException:
        raise SourceError(
            "komunikat workbook reader rejected source structure"
        ) from None
    finally:
        if book is not None:
            book.close()


def _companion_row(
    values: tuple[object, ...],
    scope: GrouperType,
    worksheet: str,
    number: int,
    *,
    ranks: bool,
) -> ProcedureRank | DepartmentSpecialty:
    try:
        if ranks:
            row = _RankRow.model_validate(
                dict(zip(("code", "title", "rank"), values, strict=True))
            )
            return ProcedureRank(
                scope, row.code, row.title, row.rank, worksheet, number
            )
        specialty = _SpecialtyRow.model_validate(
            dict(zip(("code", "service_codes", "icd_codes"), values, strict=True))
        )
        return DepartmentSpecialty(
            scope,
            specialty.code,
            specialty.service_codes,
            specialty.icd_codes,
            worksheet,
            number,
        )
    except ValidationError:
        raise SourceError("komunikat row value is unsupported") from None


def _header(
    values: tuple[object, ...], number: int, scope: GrouperType, *, ranks: bool
) -> None:
    title = (
        f"Tabela 2. Wykaz rang procedur medycznych ({scope})"
        if ranks
        else f"Tabela 1. Wykaz specjalności komórek ({scope})"
    )
    expected: tuple[object, ...] = (title, None, None)
    if number == 2:
        expected = RANK_HEADERS if ranks else SPECIALTY_HEADERS
    elif number == 3:
        expected = ("1", "2", "3") if ranks else (1, "2", 3)
    if values != expected:
        raise SourceError("komunikat header layout is unsupported")
