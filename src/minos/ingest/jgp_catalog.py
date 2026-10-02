"""Bounded attachment 1a catalog parser: group rows and base points only.

Layouts come from `jgp_catalog_layout`; nothing is located by searching for
header text. Empty catalog cells stay absent (never zero), points are stored in
exact hundredths, and repeated codes stay separate source rows for #18.
"""

import re
from collections import Counter
from collections.abc import Sequence
from typing import TYPE_CHECKING, BinaryIO, cast

from minos.domain.jgp_catalog import (
    CatalogGroup,
    CatalogScope,
    CatalogSheet,
    CatalogValue,
    ScopeMark,
    ScopeMarker,
    ValueUnit,
)
from minos.domain.reference import SourceError
from minos.ingest.jgp_catalog_layout import (
    FIRST_DATA_ROW,
    HEADER_ROWS,
    LAYOUTS,
    LEGEND,
    SCOPE_LABELS,
    CatalogLayout,
)
from minos.ingest.jgp_catalog_manifest import (
    CatalogManifest,
    CatalogSheetExpectation,
)
from minos.ingest.workbook_archive import (
    DEFAULT_WORKBOOK_LIMITS,
    WorkbookLimits,
    WorksheetLayout,
    preflight_workbook,
)

if TYPE_CHECKING:
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet
    from openpyxl.worksheet.worksheet import Worksheet

# Shared by the parser, CLI and HTTP routes; the validation edge mirrors it.
GROUP_CODE_REGEX = r"[A-Z]{1,4}[0-9]{2,3}[A-Z]{0,3}"
CODE_PATTERN = re.compile(GROUP_CODE_REGEX)
CELL_RANGE_PATTERN = re.compile(
    r"([A-Z]{1,3})([1-9][0-9]{0,4}):([A-Z]{1,3})([1-9][0-9]{0,4})"
)
PRODUCT_PATTERN = re.compile(r"5\.51\.0[12]\.[0-9]{7}")
MAX_ROW_REPEATS = 2
MARKERS: dict[int, ScopeMarker] = {1: 1, 2: 2, 3: 3}
# Generous bounds far above every reviewed value (262,424 points, 78 days)
# that keep corrupt cells out of integer storage.
MAX_POINTS_HUNDREDTHS = 100_000_000
MAX_FINANCED_DAYS = 366
type Row = tuple[object, ...]


def parse_catalog_workbook(
    source: BinaryIO,
    manifest: CatalogManifest,
    *,
    limits: WorkbookLimits = DEFAULT_WORKBOOK_LIMITS,
) -> tuple[CatalogSheet, ...]:
    try:
        import openpyxl
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError:
        raise SourceError(
            "workbook reader requires the pinned reference extra"
        ) from None
    if openpyxl.__version__ != "3.1.5":
        raise SourceError("workbook reader version is incompatible")
    preflight = preflight_workbook(source, limits=limits)
    expected = {
        LAYOUTS[sheet.layout].worksheet: (LAYOUTS[sheet.layout], sheet)
        for sheet in manifest.sheets.values()
    }
    if set(preflight) != set(expected):
        raise SourceError("catalog worksheet identities conflict")
    book = None
    try:
        book = openpyxl.load_workbook(
            source, read_only=True, data_only=False, keep_links=False
        )
        if set(book.sheetnames) != set(preflight) or not book.read_only:
            raise SourceError("workbook reader worksheet identities conflict")
        sheets: list[CatalogSheet] = []
        for title in sorted(expected):
            layout, expectation = expected[title]
            rows = _read_rows(book[title], preflight[title], layout)
            sheet = _parse_sheet(rows, layout, preflight[title])
            _check_reviewed_counts(sheet, expectation)
            sheets.append(sheet)
        _check_catalogs_disjoint(sheets)
        return tuple(sheets)
    except ValueError, KeyError, TypeError, IndexError, OSError, InvalidFileException:
        raise SourceError("catalog workbook reader rejected source structure") from None
    finally:
        if book is not None:
            book.close()


def _read_rows(
    sheet: ReadOnlyWorksheet, preflight: WorksheetLayout, layout: CatalogLayout
) -> list[Row]:
    """Read only the reviewed columns; stray cells beyond them are not data."""
    width = column_index(layout.last_column)
    sheet.reset_dimensions()
    readable = cast("Worksheet", cast(object, sheet))
    rows: list[Row] = []
    for cells in readable.iter_rows(
        max_row=preflight.rows, max_col=width, values_only=False
    ):
        if any(cell.data_type in ("f", "e", "d") for cell in cells):
            raise SourceError("catalog formula, error or date cell is unsupported")
        rows.append(tuple(cell.value for cell in cells))
    if len(rows) < HEADER_ROWS:
        raise SourceError("catalog worksheet is too short")
    # Header cells to the right of the layout would mean an added column.
    for beyond in readable.iter_rows(
        min_row=1,
        max_row=HEADER_ROWS,
        min_col=width + 1,
        max_col=max(width + 1, preflight.columns),
        values_only=True,
    ):
        if any(not _blank(value) for value in beyond):
            raise SourceError("catalog header extends beyond the reviewed layout")
    return rows


def column_index(letters: str) -> int:
    """Spreadsheet column letters to a 1-based index (A=1, CC=81)."""
    index = 0
    for letter in letters:
        index = index * 26 + ord(letter) - ord("A") + 1
    return index


def column_letters(index: int) -> str:
    letters = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(ord("A") + remainder) + letters
    return letters


def cell_range(reference: str) -> tuple[int, int, int, int]:
    """(first column, first row, last column, last row) of a merge reference."""
    match = CELL_RANGE_PATTERN.fullmatch(reference)
    if match is None:
        raise SourceError("catalog merge reference is unsupported")
    return (
        column_index(match[1]),
        int(match[2]),
        column_index(match[3]),
        int(match[4]),
    )


def _blank(value: object) -> bool:
    return value is None or value == ""


def _text(value: object) -> str:
    return " ".join(str(value).split())


def _header(rows: Sequence[Row], layout: CatalogLayout) -> str:
    first_scope = column_index(layout.first_scope_column)
    last_scope = column_index(layout.last_scope_column)
    observed: dict[str, str | int] = {}
    for row_number in range(1, HEADER_ROWS + 1):
        for column, value in enumerate(rows[row_number - 1], 1):
            if _blank(value) or (row_number == 2 and column == 1):
                continue
            if row_number == 4 and first_scope <= column <= last_scope:
                continue
            coordinate = f"{column_letters(column)}{row_number}"
            observed[coordinate] = value if type(value) is int else _text(value)
    if observed != dict(layout.header_cells):
        raise SourceError("catalog header layout is unsupported")
    title = rows[1][0]
    if not isinstance(title, str) or _text(title) not in layout.titles:
        raise SourceError("catalog title is unsupported")
    labels = tuple(_text(value) for value in rows[3][first_scope - 1 : last_scope])
    if labels != SCOPE_LABELS:
        raise SourceError("catalog scope columns are unsupported")
    return _text(title)


def _parse_sheet(
    rows: Sequence[Row], layout: CatalogLayout, preflight: WorksheetLayout
) -> CatalogSheet:
    title = _header(rows, layout)
    groups: list[CatalogGroup] = []
    chapter_rows: list[tuple[int, str]] = []
    legend_rows: list[int] = []
    for row_number, values in enumerate(rows[FIRST_DATA_ROW - 1 :], FIRST_DATA_ROW):
        if all(_blank(value) for value in values):
            continue
        first, rest_blank = values[0], all(_blank(value) for value in values[1:])
        if legend_rows or (isinstance(first, str) and _text(first) == LEGEND[0]):
            if (
                len(legend_rows) >= len(LEGEND)
                or not rest_blank
                or _text(first) != LEGEND[len(legend_rows)]
            ):
                raise SourceError("catalog footer legend is unsupported")
            legend_rows.append(row_number)
        elif isinstance(first, str) and rest_blank:
            if chapter_rows and (not groups or groups[-1].row < chapter_rows[-1][0]):
                raise SourceError(f"catalog chapter before row {row_number} is empty")
            chapter_rows.append((row_number, _text(first)))
        elif type(first) is int and chapter_rows:
            groups.append(_group(values, row_number, layout, chapter_rows[-1][1]))
        else:
            raise SourceError(f"catalog row {row_number} is unsupported")
    if len(legend_rows) != len(LEGEND):
        raise SourceError("catalog footer legend is missing")
    if not groups or groups[-1].row < chapter_rows[-1][0]:
        raise SourceError("catalog last chapter is empty")
    _check_merges(
        preflight.merges, layout, {row for row, _ in chapter_rows}, set(legend_rows)
    )
    first_scope = column_index(layout.first_scope_column)
    scopes = tuple(
        CatalogScope(position, column_letters(first_scope + position - 1), label)
        for position, label in enumerate(SCOPE_LABELS, 1)
    )
    return CatalogSheet(
        catalog=layout.catalog,
        layout=layout.version,
        worksheet=layout.worksheet,
        title=title,
        chapters=tuple(name for _, name in chapter_rows),
        scopes=scopes,
        groups=tuple(groups),
    )


def _check_merges(
    merges: Sequence[str],
    layout: CatalogLayout,
    chapter_rows: set[int],
    legend_rows: set[int],
) -> None:
    header = {merge for merge in merges if cell_range(merge)[1] <= HEADER_ROWS}
    if frozenset(header) != layout.header_merges:
        raise SourceError("catalog header merges are unsupported")
    body = set(merges) - header
    # Every reviewed chapter heading spans the full sheet width; text alone in
    # column A without that merge is not accepted as a heading.
    chapters = {f"A{row}:{layout.last_column}{row}" for row in chapter_rows}
    if not chapters <= body:
        raise SourceError("catalog chapter heading is not a merged title row")
    for merge in body - chapters:
        first_column, first_row, _, last_row = cell_range(merge)
        # The only other reviewed spans are footer legend lines.
        if first_row != last_row or first_column != 1 or first_row not in legend_rows:
            raise SourceError("catalog body merge is unsupported")


def _group(values: Row, row: int, layout: CatalogLayout, chapter: str) -> CatalogGroup:
    ordinal, code, product, name = values[0], values[1], values[2], values[3]
    if type(ordinal) is not int or ordinal < 1:
        raise SourceError(f"catalog row {row} ordinal is unsupported")
    # Reviewed sources pad some codes with spaces or NBSP ('J01AJ ',
    # '\xa05.51.01.0014052'); only surrounding whitespace is trimmed.
    if not isinstance(code, str) or CODE_PATTERN.fullmatch(code.strip()) is None:
        raise SourceError(f"catalog row {row} group code is unsupported")
    if (
        not isinstance(product, str)
        or PRODUCT_PATTERN.fullmatch(product.strip()) is None
    ):
        raise SourceError(f"catalog row {row} product code is unsupported")
    if not product.strip().startswith(layout.product_prefix):
        raise SourceError(f"catalog row {row} product belongs to another catalog")
    if not isinstance(name, str) or not name.strip():
        raise SourceError(f"catalog row {row} group name is missing")
    catalog_values = tuple(
        CatalogValue(measure.code, measure.unit, amount, measure.column)
        for measure in layout.measures
        if (amount := _amount(values, measure.column, measure.unit, row)) is not None
    )
    if not any(value.unit == "points_hundredths" for value in catalog_values):
        raise SourceError(f"catalog row {row} has no base points")
    return CatalogGroup(
        catalog=layout.catalog,
        row=row,
        ordinal=ordinal,
        code=code.strip(),
        product_code=product.strip(),
        name=name,
        procedural=name.rstrip().endswith("*"),
        chapter=chapter,
        notes=_notes(values[column_index(layout.notes_column) - 1], row),
        values=catalog_values,
        scopes=_scopes(values, layout, row),
    )


def _amount(values: Row, column: str, unit: ValueUnit, row: int) -> int | None:
    value = values[column_index(column) - 1]
    if _blank(value):
        return None
    if type(value) is int:
        amount = value * 100 if unit == "points_hundredths" else value
    elif type(value) is float and unit == "points_hundredths":
        if not -1e12 < value < 1e12:
            raise SourceError(f"catalog row {row} value exceeds its bound")
        # Workbook floats such as 12871.720000000001 must be exact hundredths.
        amount = round(value * 100)
        if abs(value * 100 - amount) > 1e-6:
            raise SourceError(f"catalog row {row} points exceed two decimals")
    else:
        raise SourceError(f"catalog row {row} value is unsupported")
    if amount <= 0:
        raise SourceError(f"catalog row {row} value is not positive")
    bound = MAX_POINTS_HUNDREDTHS if unit == "points_hundredths" else MAX_FINANCED_DAYS
    if amount > bound:
        raise SourceError(f"catalog row {row} value exceeds its bound")
    return amount


def _scopes(values: Row, layout: CatalogLayout, row: int) -> tuple[ScopeMark, ...]:
    first = column_index(layout.first_scope_column)
    marks: list[ScopeMark] = []
    for position in range(1, len(SCOPE_LABELS) + 1):
        value = values[first + position - 2]
        if _blank(value):
            continue
        marker = MARKERS.get(value) if type(value) is int else None
        if marker is None:
            raise SourceError(f"catalog row {row} scope marker is unsupported")
        marks.append(ScopeMark(position, marker))
    return tuple(marks)


def _notes(value: object, row: int) -> str | None:
    if _blank(value):
        return None
    if not isinstance(value, str):
        raise SourceError(f"catalog row {row} note is unsupported")
    # Notes distinguish the 2026 and 2027 twin rows; blank text is no note.
    return value if value.strip() else None


def _check_reviewed_counts(
    sheet: CatalogSheet, expectation: CatalogSheetExpectation
) -> None:
    codes = Counter(group.code for group in sheet.groups)
    if max(codes.values(), default=0) > MAX_ROW_REPEATS:
        raise SourceError("catalog code repeats more than twice")
    if len(sheet.groups) != expectation.expected_groups:
        raise SourceError("reviewed catalog group count conflicts")
    if len(sheet.chapters) != expectation.expected_chapters:
        raise SourceError("reviewed catalog chapter count conflicts")
    # Section URLs address chapters by source position; names must be unique.
    if len(set(sheet.chapters)) != len(sheet.chapters):
        raise SourceError("catalog chapter heading repeats")
    repeated = sum(1 for count in codes.values() if count > 1)
    if repeated != expectation.expected_repeated_codes:
        raise SourceError("reviewed repeated catalog codes conflict")
    pairs = {(group.code, group.product_code) for group in sheet.groups}
    if len({code for code, _ in pairs}) != len(pairs) or len(
        {product for _, product in pairs}
    ) != len(pairs):
        raise SourceError("catalog code and product code do not correspond")


def _check_catalogs_disjoint(sheets: Sequence[CatalogSheet]) -> None:
    seen: set[str] = set()
    for sheet in sheets:
        codes = {group.code for group in sheet.groups}
        if codes & seen:
            raise SourceError("a group code appears in two catalogs")
        seen |= codes
