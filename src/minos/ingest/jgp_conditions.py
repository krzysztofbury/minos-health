"""Bounded original/current attachment 9 extraction, with complete cell evidence."""

import re
from collections import Counter
from dataclasses import dataclass, field
from typing import TYPE_CHECKING, BinaryIO, Literal, cast

from minos.domain import jgp_conditions as c
from minos.domain.jgp_conditions_validation import validate_characteristics
from minos.domain.reference import SourceError
from minos.ingest.jgp_catalog import CODE_PATTERN
from minos.ingest.jgp_conditions_manifest import (
    ConditionsExpectation,
    ConditionsManifest,
)
from minos.ingest.workbook_archive import WorkbookLimits, preflight_workbook

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet

CONDITIONS_LIMITS = WorkbookLimits(max_rows=50_000, max_columns=9)
WARNING = "* w przypadku filtrowania należy uwzględnić wszystkie elementy zawarte w etykiecie grupy JGP"
CONTEXT_HEADER = "Kod JGP*/ listy dodatkowej/ listy ogólnej"
MARKER = re.compile(r"ICD[ -]?(9|10)(.*)")
MEMBER = re.compile(
    r"^([0-9]{2,3}(?:\.[0-9]{1,5})?|[A-Z][0-9]{2}(?:\.[0-9A-Z]{1,3})?)(?=[*+\s(])"
)
MARKS = re.compile(r"\(\*[0-9]{1,2}\)|(?<!\*)\*{1,4}(?!\*)")


@dataclass
class _List:
    identity: int
    label: str
    system: c.System
    context: str
    marker_row: int
    naming_rows: tuple[int, ...]
    members: list[c.Member] = field(default_factory=list)


def parse_conditions_workbook(
    source: BinaryIO, manifest: ConditionsManifest
) -> tuple[c.Characteristics, ...]:
    try:
        import openpyxl
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError:
        raise SourceError(
            "characteristics reader requires the pinned reference extra"
        ) from None
    if openpyxl.__version__ != "3.1.5":
        raise SourceError("characteristics reader version is incompatible")
    layouts = preflight_workbook(source, limits=CONDITIONS_LIMITS)
    expected: dict[str, Literal["1a", "1ae"]] = (
        {"Charakterystyka": "1a"}
        if manifest.layout == "nfz-9-46-2026"
        else {"Charakterystyka 1a": "1a", "Charakterystyka 1ae": "1ae"}
    )
    if set(layouts) != set(expected) or any(item.merges for item in layouts.values()):
        raise SourceError("characteristics worksheet identities or merges conflict")
    book = None
    try:
        book = openpyxl.load_workbook(
            source, read_only=True, data_only=False, keep_links=False
        )
        if set(book.sheetnames) != set(expected) or not book.read_only:
            raise SourceError("characteristics reader worksheet identities conflict")
        result: list[c.Characteristics] = []
        for name, catalog in sorted(expected.items()):
            readable = cast("Worksheet", cast(object, book[name]))
            book[name].reset_dimensions()
            rows: list[tuple[str | None, ...]] = []
            for cells in readable.iter_rows(
                max_row=layouts[name].rows, max_col=layouts[name].columns
            ):
                if any(
                    cell.data_type in ("f", "e", "d")
                    or (cell.value is not None and not isinstance(cell.value, str))
                    for cell in cells
                ):
                    raise SourceError(
                        "characteristics formula, error, date or non-text cell"
                    )
                values = tuple(cast(str | None, cell.value) for cell in cells)
                if any(value and value.strip() for value in values[3:]):
                    raise SourceError(
                        "characteristics nonempty column outside reviewed layout"
                    )
                if any(value is not None and len(value) > 8192 for value in values[:3]):
                    raise SourceError("characteristics cell exceeds text bound")
                rows.append(values[:3])
            start = _headers(rows, catalog, manifest.layout)
            sheet = _extract(rows, start, catalog, name)
            from minos.ingest.jgp_condition_phrases import interpret_sheet

            sheet = interpret_sheet(sheet)
            validate_characteristics(sheet)
            check_counts(sheet, manifest.sheets[catalog])
            result.append(sheet)
        return tuple(result)
    except (
        ValueError,
        KeyError,
        TypeError,
        IndexError,
        OSError,
        InvalidFileException,
    ) as error:
        raise SourceError(
            "characteristics workbook structure is unsupported"
        ) from error
    finally:
        if book is not None:
            book.close()


def _headers(rows: list[tuple[str | None, ...]], catalog: str, layout: str) -> int:
    header = 2 if layout == "nfz-9-46-2026" else 3
    if len(rows) < header + 2:
        raise SourceError("characteristics sheet is truncated")
    if header == 2 and rows[0] != (WARNING, None, None):
        raise SourceError("original characteristics filter warning changed")
    if header == 3 and any(value for value in rows[0]):
        raise SourceError("characteristics leading row changed")
    if header == 3 and rows[1] != (WARNING, None, None):
        raise SourceError("characteristics filter warning changed")
    observed = tuple(" ".join((value or "").split()) for value in rows[header - 1])
    title = "Charakterystyka JGP" if header == 2 else "Charakterystyka JGP " + catalog
    if observed != ("Sekcja", CONTEXT_HEADER, title):
        raise SourceError("characteristics header layout is unsupported")
    return header + 1


def _extract(
    rows: list[tuple[str | None, ...]],
    start: int,
    catalog: Literal["1a", "1ae"],
    name: str,
) -> c.Characteristics:
    evidence: list[c.Evidence] = []
    groups: list[c.Group] = []
    lists: list[_List] = []
    active: _List | None = None
    heading: tuple[str, int] | None = None
    previous_context: str | None = None
    for number, values in enumerate(rows[start - 1 :], start):
        section, context, text = values
        if not text or not text.strip():
            if section or context:
                evidence.append(
                    c.Evidence(number, section, context, text or "", "section")
                )
            continue
        stripped, key = text.strip(), (context or "").strip()
        if key != previous_context:
            active, heading = None, None
            previous_context = key
        marker = MARKER.fullmatch(stripped)
        code = MEMBER.match(stripped)
        category: c.EvidenceCategory
        if not key:
            category = "section"
            active, heading = None, None
        elif CODE_PATTERN.fullmatch(key) and stripped.startswith(key + " "):
            category = "title"
            groups.append(c.Group(key, number, ()))
            active, heading = None, None
        elif marker:
            category = "marker"
            system: c.System = "icd9-pl" if marker[1] == "9" else "icd10-pl"
            label, naming = _list_name(key, marker[2], heading, evidence, system)
            active = _List(len(lists) + 1, label, system, key, number, naming)
            lists.append(active)
            heading = None
        elif active is not None and code:
            category = "member"
            active.members.append(
                c.Member(number, code[1], tuple(MARKS.findall(stripped)))
            )
        elif code:
            category = "unbound_member"
        elif stripped == "lub":
            category = "connector"
            active = None
        elif re.fullmatch(r"[A-Z]{1,4}[0-9]{1,3}[A-Za-z]{0,3}|Lista Ogólna", stripped):
            category = "list_heading"
            heading = (stripped, number)
            active = None
        else:
            category = "label"
            active = None
        evidence.append(c.Evidence(number, section, context, text, category))
    if len({g.code for g in groups}) != len(groups):
        raise SourceError("characteristics group title is duplicated")
    frozen = tuple(
        c.DirectionalList(
            item.identity,
            item.label,
            item.system,
            item.context,
            item.marker_row,
            item.naming_rows,
            tuple(item.members),
        )
        for item in lists
    )
    return c.Characteristics(
        catalog, name, tuple(evidence), tuple(groups), frozen, (), ()
    )


def _list_name(
    context: str,
    suffix: str,
    heading: tuple[str, int] | None,
    evidence: list[c.Evidence],
    system: c.System,
) -> tuple[str, tuple[int, ...]]:
    suffix = suffix.strip().strip(":")
    if suffix.startswith("(") and suffix.endswith(")"):
        suffix = suffix[1:-1].strip()
    if heading is not None:
        return heading[0], (heading[1],)
    if suffix:
        # Explicit marker names (ICD 9 N03a) differ from suffix markers (ICD-9a).
        return (suffix if re.search(r"[0-9]", suffix) else context + suffix), ()
    # The source explicitly names a contextual list as "poniższej" in A03.
    for item in reversed(evidence[-16:]):
        if (item.context or "").strip() != context:
            break
        previous_marker = MARKER.fullmatch(item.text.strip())
        if (
            previous_marker is not None
            and ("icd9-pl" if previous_marker[1] == "9" else "icd10-pl") == system
        ):
            break
        below = re.search(
            r"rozpoznania (?:zasadniczego|współistniejącego) z poniższej listy"
            + r"(?:\s+\(([A-Z]{1,4}[0-9]{1,3}[A-Za-z]{0,3})\))?",
            item.text,
        )
        if below is not None and system == "icd10-pl":
            # A source-local unnamed declaration must not alias a named list
            # elsewhere in the same group. Its row supplies physical identity.
            return below[1] or "poniższa lista", (item.row,)
        matches = re.findall(
            r"z listy (?:procedur |rozpoznań )?([A-Z][A-Za-z0-9]+) \(poniższej\)",
            item.text,
        )
        if matches and system == "icd9-pl":
            if len(set(matches)) == 1:
                return matches[0], (item.row,)
    return context, ()


def condition_counts(sheet: c.Characteristics) -> dict[str, object]:
    return {
        "groups": len(sheet.groups),
        "evidence_cells": len(sheet.evidence),
        "unbound_member_occurrences": sum(
            e.category == "unbound_member" for e in sheet.evidence
        ),
        "lists": len(sheet.lists),
        "member_occurrences": sum(len(item.members) for item in sheet.lists),
        "unique_members": sum(
            len({m.code for m in item.members}) for item in sheet.lists
        ),
        "conditions": len(sheet.conditions),
        "kind_counts": dict(
            sorted(Counter(n.kind or "unsupported" for n in sheet.conditions).items())
        ),
        "unsupported_reasons": dict(
            sorted(Counter(n.reason for n in sheet.conditions if n.reason).items())
        ),
    }


def check_counts(sheet: c.Characteristics, expected: ConditionsExpectation) -> None:
    observed = condition_counts(sheet)
    reviewed = expected.model_dump(exclude={"unmatched_codes"})
    if observed != reviewed:
        raise SourceError("reviewed characteristics counts do not match")


def validate_source_projection(sheet: c.Characteristics) -> None:
    """Reconstruct extraction/interpretation from persisted A/B/C evidence.

    Counts and foreign keys cannot detect a changed source label, physical list
    name, or dropped unknown requirement when a checksum was recomputed.
    """
    validate_characteristics(sheet)
    rows: list[tuple[str | None, ...]] = [(None, None, None)] * max(
        e.row for e in sheet.evidence
    )
    for item in sheet.evidence:
        rows[item.row - 1] = (item.section, item.context, item.text)
    extracted = _extract(
        rows, min(e.row for e in sheet.evidence), sheet.catalog, sheet.worksheet
    )
    from minos.ingest.jgp_condition_phrases import interpret_sheet

    if interpret_sheet(extracted) != sheet:
        raise SourceError("characteristics projection conflicts with source evidence")
