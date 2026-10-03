import json
from datetime import date
from pathlib import Path
from typing import TYPE_CHECKING, cast

import pytest
from fastapi.testclient import TestClient

from minos.domain.jgp_ruleset import NoRulesetForDate
from minos.http.app import create_app
from minos.ingest.reference_artifact import build_reference_artifact
from minos.ingest.reference_manifest import DatasetManifest
from minos.store.jgp_condition_reference import (
    code_memberships,
    group_conditions,
    list_members,
)
from minos.store.reference import open_reference_bundle
from tests.fixtures.jgp_catalog import catalog_workbook, default_sheets, stage_catalog
from tests.fixtures.jgp_conditions import manifest as source_manifest
from tests.fixtures.jgp_conditions import workbook
from tests.integration.test_jgp_conditions_reference import conditions_manifest

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet


def consumer_bundle(
    root: Path,
    *,
    unsupported: bool = False,
    long_list: bool = False,
    title_only: bool = False,
    many_groups: bool = False,
) -> Path:
    base = conditions_manifest(root)
    characteristics = base.characteristics[0]
    if unsupported or long_list or title_only or many_groups:

        def edit(sheet: Worksheet) -> None:
            if sheet.title != "Charakterystyka 1a":
                return
            if unsupported:
                sheet.append(("A", "A01", "Unreviewed allocation requirement"))
            if title_only:
                sheet["C5"] = None
            if many_groups:
                for number in range(101):
                    code = f"X{number:03}"
                    sheet.append(("X", code, f"{code} Synthetic reverse group"))
                    sheet.append(
                        (
                            "X",
                            code,
                            "wymagane wskazanie rozpoznania zasadniczego z Listy Ogólnej",
                        )
                    )
            if long_list:
                sheet.insert_rows(13, 150)
                for row in range(13, 163):
                    _ = sheet.cell(row, 1, "A")
                    _ = sheet.cell(row, 2, "A01")
                    _ = sheet.cell(row, 3, "E11.0 Duplicate source occurrence")

        content = workbook(edit=edit)
        characteristics = source_manifest(content)
        expected = characteristics.sheets["1a"].model_dump()
        expected["evidence_cells"] += (1 if unsupported else 0) + (
            150 if long_list else 0
        )
        if unsupported:
            expected["conditions"] += 1
            expected["kind_counts"]["unsupported"] = 1
            expected["unsupported_reasons"] = {"unconsumed_source_text": 1}
        if long_list:
            expected["member_occurrences"] += 150
        if title_only:
            expected["conditions"] -= 2
            expected["kind_counts"] = {"program_membership": 2, "unsupported": 1}
            expected["unsupported_reasons"] = {"missing_label": 1}
        if many_groups:
            expected["groups"] += 101
            expected["evidence_cells"] += 202
            expected["conditions"] += 101
            expected["kind_counts"]["diagnosis_presence"] += 101
        from minos.ingest.jgp_conditions_manifest import ConditionsExpectation

        characteristics = characteristics.model_copy(
            update={
                "sheets": {
                    **characteristics.sheets,
                    "1a": ConditionsExpectation.model_validate(expected),
                }
            }
        )
        _ = (root / "blobs/sha256" / characteristics.workbook.sha256).write_bytes(
            content
        )
    catalog = stage_catalog(
        root,
        catalog_workbook(default_sheets()),
        release={
            "baw_document_id": 44042,
            "document_url": "https://baw.nfz.gov.pl/NFZ/document/44042/Synthetic",
        },
    )
    data = base.model_dump(mode="json")
    data["characteristics"] = [
        characteristics.model_copy(
            update={
                "catalog_release_id": catalog.release_id,
            }
        ).model_dump(mode="json")
    ]
    data.update(
        catalogs=[catalog.model_dump(mode="json")],
        catalog_knowledge_as_of="2026-10-02",
        catalog_knowledge_evidence="Synthetic consumer review",
    )
    selected = DatasetManifest.model_validate_json(json.dumps(data))
    _ = build_reference_artifact(selected, root, root / "bundle")
    return root / "bundle"


def test_dated_reference_members_keep_duplicates_exact_codes_and_missing_correspondence(
    tmp_path: Path,
) -> None:
    bundle = consumer_bundle(tmp_path)
    with open_reference_bundle(bundle) as (connection, _):
        day = date(2026, 7, 1)
        group = group_conditions(connection, "1a", "A01", day)
        assert group["status"] == "fully_interpreted"
        assert len(cast(list[object], group["nodes"])) == 3
        assert cast(dict[str, object], group["source"])["catalog"] == "1a"
        members = list_members(connection, "1a", 2, day)
        entries = cast(list[dict[str, object]], members["members"])
        assert [(e["code"], e["row"], e["correspondence"]) for e in entries] == [
            ("G60.4", 10, "not_in_reference"),
            ("G60.4", 11, "not_in_reference"),
            ("E11.0", 12, "matched"),
        ]
        reverse = code_memberships(connection, "icd10-pl", "G60.4", "1a", day)
        assert reverse["total"] == 1
        assert [
            (g["code"], g["list_ids"])
            for g in cast(list[dict[str, object]], reverse["groups"])
        ] == [("A01", [2])]
        assert code_memberships(connection, "icd10-pl", "G60", "1a", day)["total"] == 0
        assert code_memberships(connection, "icd9-pl", "G60.4", "1a", day)["total"] == 0
        with pytest.raises(NoRulesetForDate):
            _ = group_conditions(connection, "1a", "A01", date(2026, 6, 30))


def test_http_exposes_source_traced_rules_and_dated_list_pages(tmp_path: Path) -> None:
    with TestClient(create_app(consumer_bundle(tmp_path))) as client:
        response = client.get("/api/v1/jgp/2026-07-01/groups/A01")
        assert response.status_code == 200
        assert response.json()["conditions"]["status"] == "fully_interpreted"
        listing = client.get("/api/v1/jgp/2026-07-01/lists/1a/2/pages/1")
        assert listing.status_code == 200
        assert listing.json()["total_occurrences"] == 3
        assert (
            client.get("/api/v1/jgp/2026-07-01/lists/1a/2/pages/2").status_code == 404
        )
        assert (
            client.get("/api/v1/jgp/2026-06-30/lists/1a/2/pages/1").status_code == 404
        )
        assert (
            client.get("/api/v1/jgp/2026-07-01/lists/1a/2/pages/1?x=1").status_code
            == 400
        )


def test_one_unsupported_root_hides_every_structured_rule_but_not_source_membership(
    tmp_path: Path,
) -> None:
    with TestClient(create_app(consumer_bundle(tmp_path, unsupported=True))) as client:
        answer = cast(
            dict[str, object],
            client.get("/api/v1/jgp/2026-07-01/groups/A01").json()["conditions"],
        )
        assert answer["status"] == "not_yet_interpreted"
        assert answer["nodes"] == answer["selectors"] == answer["roots"] == []
        assert any(
            "Unreviewed allocation" in cast(str, row["text"])
            for row in cast(list[dict[str, object]], answer["source_text"])
        )
        page = client.get("/jgp/2026-07-01/A01")
        assert page.status_code == 200
        assert "Jeszcze nie zinterpretowano całej grupy" in page.text
        assert 'class="jgp-rules"' not in page.text
        concept = client.get("/slowniki/icd-10/E11.0/jgp/2026-07-01")
        assert concept.status_code == 200
        assert 'href="/jgp/2026-07-01/A01"' in concept.text
        assert "jeszcze nie zinterpretowano" in concept.text


def test_list_pagination_is_physical_and_missing_dictionary_codes_are_not_links(
    tmp_path: Path,
) -> None:
    with TestClient(create_app(consumer_bundle(tmp_path, long_list=True))) as client:
        first = cast(
            dict[str, object],
            client.get("/api/v1/jgp/2026-07-01/lists/1a/2/pages/1").json(),
        )
        second = cast(
            dict[str, object],
            client.get("/api/v1/jgp/2026-07-01/lists/1a/2/pages/2").json(),
        )
        assert first["total_occurrences"] == 153
        first_members = cast(list[dict[str, object]], first["members"])
        second_members = cast(list[dict[str, object]], second["members"])
        assert len(first_members) == 100 and len(second_members) == 53
        assert cast(int, first_members[-1]["row"]) < cast(int, second_members[0]["row"])
        page = client.get("/jgp/2026-07-01/listy/1a/2/1")
        assert page.status_code == 200
        assert 'href="/slowniki/icd-10/G60.4' not in page.text
        assert 'href="/slowniki/icd-10/E11.0/jgp/2026-07-01"' in page.text
        assert 'href="/jgp/2026-07-01/listy/1a/2/2"' in page.text
        assert client.get("/jgp/2026-07-01/listy/1a/2/3").status_code == 404


def test_conditions_only_group_has_no_invented_product_or_points(
    tmp_path: Path,
) -> None:
    with TestClient(create_app(consumer_bundle(tmp_path))) as client:
        answer = client.get("/api/v1/jgp/2026-07-01/groups/N09")
        assert answer.status_code == 200
        assert answer.json()["group"]["product_code"] is None
        assert answer.json()["base_values"] == []
        page = client.get("/jgp/2026-07-01/N09")
        assert page.status_code == 200
        assert "bez wiersza katalogu na ten dzień" in page.text
        assert "Nie oznacza to wartości zero" in page.text


def test_group_without_requirement_text_stays_uninterpreted_in_reverse_membership(
    tmp_path: Path,
) -> None:
    bundle = consumer_bundle(tmp_path, title_only=True)
    with open_reference_bundle(bundle) as (connection, _):
        day = date(2026, 7, 1)
        assert (
            group_conditions(connection, "1a", "A01", day)["status"]
            == "not_yet_interpreted"
        )
        reverse = code_memberships(connection, "icd10-pl", "E11.0", "1a", day)
        assert reverse["total"] == 1
        assert (
            cast(list[dict[str, object]], reverse["groups"])[0]["fully_interpreted"]
            is False
        )


def test_reverse_pages_keep_one_group_each_and_dated_catalog_links(
    tmp_path: Path,
) -> None:
    bundle = consumer_bundle(tmp_path, many_groups=True)
    with open_reference_bundle(bundle) as (connection, _):
        pages = [
            code_memberships(
                connection, "icd10-pl", "E11.0", "1a", date(2026, 7, 1), page
            )
            for page in (1, 2)
        ]
        assert [page["total"] for page in pages] == [102, 102]
        codes = [
            [g["code"] for g in cast(list[dict[str, object]], page["groups"])]
            for page in pages
        ]
        assert len(codes[0]) == 100 and codes[1] == ["X099", "X100"]
        assert len(set(codes[0] + codes[1])) == 102
    with TestClient(create_app(bundle)) as client:
        first = client.get("/slowniki/icd-10/E11.0/jgp/2026-07-01")
        assert 'href="/slowniki/icd-10/E11.0/jgp/2026-07-01/1a/2"' in first.text
        second = client.get("/slowniki/icd-10/E11.0/jgp/2026-07-01/1a/2")
        assert second.status_code == 200
        assert 'href="/jgp/2026-07-01/X100"' in second.text
        assert 'href="/jgp/2026-07-01/A01"' not in second.text
        assert (
            client.get("/slowniki/icd-10/E11.0/jgp/2026-07-01/1a/3").status_code == 404
        )


def test_all_condition_consumers_switch_to_the_governing_release_at_date_boundary(
    tmp_path: Path,
) -> None:
    from minos.ingest.jgp_conditions_manifest import ConditionsExpectation
    from minos.store.reference import read_artifact_manifest

    bundle = consumer_bundle(tmp_path)
    dataset = read_artifact_manifest(bundle).dataset
    current_catalog = dataset.catalogs[0]
    previous_catalog = stage_catalog(
        tmp_path,
        catalog_workbook(default_sheets()),
        release={
            "release_id": "catalog-previous",
            "baw_document_id": 43868,
            "document_url": "https://baw.nfz.gov.pl/NFZ/document/43868/Synthetic",
            "service_from": date(2026, 5, 1),
            "published_on": date(2026, 4, 30),
            "act_date": date(2026, 4, 29),
            "entry_into_force": date(2026, 5, 1),
        },
    )
    current_catalog = current_catalog.model_copy(
        update={"supersedes": previous_catalog.release_id}
    )

    def previous_members(sheet: Worksheet) -> None:
        if sheet.title == "Charakterystyka":
            sheet["C11"] = "G60.4 Previous-only literal diagnosis"

    content = workbook(original=True, edit=previous_members)
    previous = source_manifest(content, original=True)
    counts = previous.sheets["1a"].model_dump()
    counts["unique_members"] -= 1
    previous = previous.model_copy(
        update={
            "catalog_release_id": previous_catalog.release_id,
            "sheets": {"1a": ConditionsExpectation.model_validate(counts)},
        }
    )
    _ = (tmp_path / "blobs/sha256" / previous.workbook.sha256).write_bytes(content)
    data = dataset.model_dump(mode="json")
    data["catalogs"] = [
        previous_catalog.model_dump(mode="json"),
        current_catalog.model_dump(mode="json"),
    ]
    data["characteristics"] = [
        previous.model_dump(mode="json"),
        dataset.characteristics[0]
        .model_copy(update={"supersedes": previous.release_id})
        .model_dump(mode="json"),
    ]
    selected = DatasetManifest.model_validate_json(json.dumps(data))
    _ = build_reference_artifact(selected, tmp_path, tmp_path / "two-releases")
    with TestClient(create_app(tmp_path / "two-releases")) as client:
        for date_text, release_id, has_e11 in (
            ("2026-06-30", previous.release_id, False),
            ("2026-07-01", dataset.characteristics[0].release_id, True),
        ):
            group = cast(
                dict[str, object],
                client.get(f"/api/v1/jgp/{date_text}/groups/A01").json(),
            )
            conditions = cast(dict[str, object], group["conditions"])
            assert (
                cast(dict[str, object], conditions["source"])["release_id"]
                == release_id
            )
            listing = cast(
                dict[str, object],
                client.get(f"/api/v1/jgp/{date_text}/lists/1a/2/pages/1").json(),
            )
            assert (
                cast(dict[str, object], listing["source"])["release_id"] == release_id
            )
            assert (
                any(
                    m["code"] == "E11.0"
                    for m in cast(list[dict[str, object]], listing["members"])
                )
                is has_e11
            )
            reverse = client.get(f"/slowniki/icd-10/E11.0/jgp/{date_text}")
            assert reverse.status_code == 200 and release_id in reverse.text
            assert (f'href="/jgp/{date_text}/A01"' in reverse.text) is has_e11
            page = client.get(f"/jgp/{date_text}/listy/1a/2/1")
            assert (
                f'href="/slowniki/icd-10/E11.0/jgp/{date_text}"' in page.text
            ) is has_e11


def test_dated_concept_routes_keep_dates_errors_cache_and_query_privacy(
    tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    import logging

    with TestClient(create_app(consumer_bundle(tmp_path))) as client:
        caplog.set_level(logging.INFO, logger="minos.http")
        dated = client.get("/slowniki/icd-10/E11.0/jgp/2026-07-01")
        assert dated.status_code == 200
        assert dated.headers["cache-control"].startswith("public")
        assert 'content="noindex, follow"' in dated.text
        undated = client.get("/slowniki/icd-10/E11.0")
        assert (
            undated.status_code == 200
            and undated.headers["cache-control"] == "no-store"
        )
        assert (
            '<link rel="canonical" href="https://minoshealth.pl/slowniki/icd-10/E11.0/jgp/2026-07-01">'
            in dated.text
        )
        assert 'href="/jgp/2026-07-01/A01"' in dated.text
        for day in ("2026-06-30", "20260701", "2026-02-30"):
            response = client.get(f"/slowniki/icd-10/E11.0/jgp/{day}")
            assert (
                response.status_code == 404
                and response.headers["cache-control"] == "no-store"
            )
        refused = client.get(
            "/slowniki/icd-10/E11.0/jgp/2026-07-01?query=PRIVATE_QUERY_SENTINEL"
        )
        assert refused.status_code == 400
        assert (
            "PRIVATE_QUERY_SENTINEL" not in caplog.text and "E11.0" not in caplog.text
        )


def test_pending_modifier_candidate_never_changes_catalog_binding(
    tmp_path: Path,
) -> None:
    from minos.store.reference import read_artifact_manifest
    from tests.integration.test_jgp_modifiers_reference import stage_modifiers

    bundle = consumer_bundle(tmp_path)
    dataset = read_artifact_manifest(bundle).dataset
    data = dataset.model_dump(mode="json")
    modifier = stage_modifiers(tmp_path)
    data["modifiers"] = [modifier.model_dump(mode="json")]
    catalog_id = dataset.catalogs[0].release_id
    data["modifier_candidates"] = [
        {
            "catalog_release_id": catalog_id,
            "modifier_release_id": modifier.release_id,
            "status": "pending_legal_review",
            "evidence": "Synthetic pending review; unknown group A02 is retained",
        }
    ]
    selected = DatasetManifest.model_validate_json(json.dumps(data))
    _ = build_reference_artifact(selected, tmp_path, tmp_path / "candidate-bundle")
    with TestClient(create_app(tmp_path / "candidate-bundle")) as client:
        answer = cast(
            dict[str, object], client.get("/api/v1/jgp/2026-07-01/groups/A01").json()
        )
        ruleset = cast(dict[str, object], answer["ruleset"])
        modifiers = cast(dict[str, object], ruleset["modifiers"])
        assert modifiers["binding"] == "unbound"
        candidate = cast(dict[str, object], modifiers["review_candidate"])
        assert candidate["generation"] == 100
        assert candidate["status"] == "pending_legal_review"
        assert "final_points" not in answer
    report = (tmp_path / "candidate-bundle/review.txt").read_text()
    assert "A02" in report and "Absence blocks complete binding" in report
