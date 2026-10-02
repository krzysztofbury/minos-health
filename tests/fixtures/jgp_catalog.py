"""Synthetic attachment 1a workbooks shaped like the reviewed BAW layouts.

Headers come from the layout constants, so these fixtures exercise parser
behavior; the constants themselves are proven against the pinned BAW bytes by
tests/compatibility/test_jgp_catalog_private_capture.py.
"""

import hashlib
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from typing import TYPE_CHECKING, cast

from minos.ingest.jgp_catalog import column_index, column_letters
from minos.ingest.jgp_catalog_layout import (
    LAYOUTS,
    LEGEND,
    SCOPE_LABELS,
    CatalogLayout,
)
from minos.ingest.jgp_catalog_manifest import CatalogManifest

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet


@dataclass(frozen=True)
class Row:
    ordinal: int
    code: object
    product: object
    name: object
    values: dict[str, object] = field(default_factory=dict[str, object])
    scopes: dict[int, object] = field(default_factory=dict[int, object])
    notes: object = None


ROWS_1A = (
    Row(1, "A01", "5.51.01.0001001", "Zabiegi wewnątrzczaszkowe *",
        {"E": 13586, "F": "", "H": 25, "L": 641}, {15: 1, 50: 3}),
    Row(2, "PZD02", "\xa05.51.01.0018028", "Duże zabiegi klatki piersiowej *",
        {"E": 13991, "F": 12871.720000000001, "H": 21}, {9: 2}),
    Row(3, "H21", "5.51.01.0008021", "Artroskopowa rekonstrukcja *",
        {"E": 3974, "I": 3657}, {15: 1, 58: 1},
        "- w zakresie chirurgia ogólna możliwość rozliczania do 31.12.2026 r"),
    Row(4, "H21", "5.51.01.0008021", "Artroskopowa rekonstrukcja *",
        {"E": 3974, "I": 3657}, {58: 1},
        "-dotyczy świadczeń sprawozdanych i\xa0 rozliczanych od 1.01.2027 r."),
    Row(5, "J01AJ ", "5.51.01.0010001", "Choroby zachowawcze", {"E": 2170, "G": 1200}),
)  # fmt: skip
ROWS_1AE = (
    Row(1, "AU35C", "5.51.02.0001001", "Choroby zwyrodnieniowe OUN -AE",
        {"E": 3526, "H": 3526, "I": 19, "J": 559}, {25: 1, 49: 3}),
    Row(2, "SU52", "5.51.02.0016005", "Niedobory odporności -AE",
        {"E": 1446, "H": 1446}, {1: 3}),
)  # fmt: skip


def catalog_workbook(
    sheets: dict[str, tuple[Row, ...]],
    *,
    shift_from: str | None = None,
    edit: Callable[[Worksheet, CatalogLayout], None] | None = None,
) -> bytes:
    """Build reviewed-layout sheets; `shift_from` inserts a column before it."""
    from openpyxl import Workbook

    book = Workbook()
    active = book.active
    assert active is not None
    book.remove(active)
    for version, rows in sheets.items():
        layout = LAYOUTS[version]
        sheet = cast("Worksheet", book.create_sheet(layout.worksheet))
        place = _placer(shift_from)
        last = place(layout.last_column)
        sheet.merge_cells(f"A1:{last}1")
        sheet.merge_cells(f"A2:{last}2")
        sheet["A2"] = sorted(layout.titles)[-1]
        for coordinate, header in layout.header_cells:
            letters = coordinate.rstrip("0123456789")
            sheet[place(letters) + coordinate[len(letters) :]] = header
        first_scope = column_index(place(layout.first_scope_column))
        for offset, label in enumerate(SCOPE_LABELS):
            _ = sheet.cell(4, first_scope + offset, label)
        for merge in sorted(layout.header_merges):
            if merge in (f"A1:{layout.last_column}1", f"A2:{layout.last_column}2"):
                continue
            start, end = merge.split(":")
            sheet.merge_cells(f"{_shift(start, place)}:{_shift(end, place)}")
        sheet["A6"] = "Choroby układu nerwowego"
        sheet.merge_cells(f"A6:{last}6")
        for number, row in enumerate(rows, 7):
            _write_row(sheet, number, row, layout, place)
        footer = 8 + len(rows)
        for line_number, line in enumerate(LEGEND, footer):
            _ = sheet.cell(line_number, 1, line)
        # Real workbooks keep stray formulas far right of the layout.
        _ = sheet.cell(7, column_index(last) + 30, "=#REF!")
        if edit is not None:
            edit(sheet, layout)
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def _placer(shift_from: str | None) -> Callable[[str], str]:
    threshold = column_index(shift_from) if shift_from else 10_000

    def place(letters: str) -> str:
        index = column_index(letters)
        return column_letters(index + 1 if index >= threshold else index)

    return place


def _shift(coordinate: str, place: Callable[[str], str]) -> str:
    column = coordinate.rstrip("0123456789")
    return place(column) + coordinate[len(column) :]


def _write_row(
    sheet: Worksheet,
    number: int,
    row: Row,
    layout: CatalogLayout,
    place: Callable[[str], str],
) -> None:
    for column, value in (
        ("A", row.ordinal),
        ("B", row.code),
        ("C", row.product),
        ("D", row.name),
    ):
        sheet[f"{place(column)}{number}"] = cast(str, value)
    for column, value in row.values.items():
        sheet[f"{place(column)}{number}"] = cast(str, value)
    first_scope = column_index(place(layout.first_scope_column))
    for position, value in row.scopes.items():
        _ = sheet.cell(number, first_scope + position - 1, cast(str, value))
    if row.notes is not None:
        sheet[f"{place(layout.notes_column)}{number}"] = cast(str, row.notes)


def default_sheets() -> dict[str, tuple[Row, ...]]:
    return {"nfz-1a-74-2026": ROWS_1A, "nfz-1ae-74-2026": ROWS_1AE}


def _repeated(rows: tuple[Row, ...]) -> int:
    codes = [str(row.code).strip() for row in rows]
    return len({code for code in codes if codes.count(code) > 1})


def stage_catalog(
    root: Path,
    content: bytes,
    *,
    sheets: dict[str, tuple[Row, ...]] | None = None,
    repeated: dict[str, int] | None = None,
    release: dict[str, object] | None = None,
) -> CatalogManifest:
    chosen = sheets if sheets is not None else default_sheets()
    archive = root / "blobs/sha256"
    archive.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256(content).hexdigest()
    _ = (archive / digest).write_bytes(content)
    expectations = {
        LAYOUTS[version].catalog: {
            "layout": version,
            "expected_groups": len(rows),
            "expected_chapters": 1,
            "expected_repeated_codes": (repeated or {}).get(
                LAYOUTS[version].catalog, _repeated(rows)
            ),
        }
        for version, rows in chosen.items()
    }
    return CatalogManifest.model_validate(
        {
            "release_id": "jgp-catalog-synthetic",
            "act_number": "46/2026/DSOZ",
            "baw_document_id": 43943,
            "automatic_type": 3,
            "source_document_id": 43868,
            "introductory_document_id": 43942,
            "act_date": datetime(2026, 7, 23).date(),
            "entry_into_force": datetime(2026, 7, 24).date(),
            "published_on": datetime(2026, 7, 23).date(),
            "service_from": datetime(2026, 7, 1).date(),
            "service_evidence": "Synthetic: services from 2026-07-01.",
            "supersedes": None,
            "row_applicability": _twin_windows(chosen),
            "modifier_binding": "unbound",
            "document_url": "https://baw.nfz.gov.pl/NFZ/document/43943/Synthetic",
            "publisher": "Synthetic NFZ fixture",
            "license": "Synthetic fixture, not publisher source bytes",
            "notices": ("Catalog base points without modifiers.",),
            "workbook": {
                "source_url": "https://baw.nfz.gov.pl/api/file/4/30983",
                "name": "Zalacznik1a.xlsx",
                "sha256": digest,
                "size_bytes": len(content),
                "fetched_at": datetime(2026, 10, 2, tzinfo=UTC),
            },
            "sheets": expectations,
            **(release or {}),
        },
        strict=False,
    )


def _twin_windows(sheets: dict[str, tuple[Row, ...]]) -> tuple[dict[str, object], ...]:
    """Reviewed windows for the synthetic H21 twins, as 102/2026 § 4 ust. 5."""
    rows = sheets.get("nfz-1a-74-2026", ())
    if [row.ordinal for row in rows if row.code == "H21"] != [3, 4]:
        return ()
    return (
        {"catalog": "1a", "ordinals": (4,), "service_from": datetime(2027, 1, 1).date(),
         "service_to": None, "note_marker": "od 1.01.2027", "evidence": "Synthetic § 4 ust. 5."},
        {"catalog": "1a", "ordinals": (3,), "service_from": None,
         "service_to": datetime(2026, 12, 31).date(), "note_marker": "do 31.12.2026",
         "evidence": "Synthetic twin until 2026-12-31."},
    )  # fmt: skip
