from collections.abc import Callable
from dataclasses import replace
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import TYPE_CHECKING
from zipfile import ZipFile

import pytest
from pydantic import ValidationError

from minos.domain.jgp_catalog import CatalogGroup, CatalogValue, ScopeMark
from minos.domain.reference import SourceError
from minos.ingest.jgp_catalog import column_index, parse_catalog_workbook
from minos.ingest.jgp_catalog_layout import (
    LEGEND,
    NFZ_1A_74_2026,
    CatalogLayout,
    check_layout,
)
from minos.ingest.jgp_catalog_manifest import CatalogManifest
from tests.fixtures.jgp_catalog import (
    ROWS_1A,
    ROWS_1AE,
    Row,
    catalog_workbook,
    default_sheets,
    stage_catalog,
)
from tests.fixtures.komunikat import package_bytes

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet


def parse(
    tmp_path: Path, content: bytes, **staging: dict[str, int]
) -> tuple[CatalogGroup, ...]:
    manifest = stage_catalog(tmp_path, content, repeated=staging.get("repeated"))
    sheets = parse_catalog_workbook(BytesIO(content), manifest)
    return tuple(group for sheet in sheets for group in sheet.groups)


def test_catalog_rows_keep_exact_base_points_absence_and_source_rows(
    tmp_path: Path,
) -> None:
    groups = parse(tmp_path, catalog_workbook(default_sheets()))
    first = groups[0]
    assert first == CatalogGroup(
        catalog="1a",
        row=7,
        ordinal=1,
        code="A01",
        product_code="5.51.01.0001001",
        name="Zabiegi wewnątrzczaszkowe *",
        procedural=True,
        chapter="Choroby układu nerwowego",
        notes=None,
        values=(
            CatalogValue("hospitalization", "points_hundredths", 1358600, "E"),
            CatalogValue("financed_stay", "days", 25, "H"),
            CatalogValue("extra_stay_day", "points_hundredths", 64100, "L"),
        ),
        scopes=(ScopeMark(15, 1), ScopeMark(50, 3)),
    )
    planned = groups[1].values[1]
    assert (planned.measure, planned.amount) == ("planned_hospitalization", 1287172)
    assert groups[1].product_code == "5.51.01.0018028"
    repeated = [group for group in groups if group.code == "H21"]
    assert [(group.row, len(group.scopes)) for group in repeated] == [(9, 2), (10, 1)]
    assert repeated[1].notes == ROWS_1A[3].notes
    assert groups[4].code == "J01AJ" and groups[4].procedural is False
    day_stay = [value.measure for value in groups[5].values]
    assert (groups[5].catalog, day_stay) == (
        "1ae",
        [
            "hospitalization",
            "day_stay_under_12_hours",
            "financed_stay_excluding_day_stay",
            "extra_stay_day_excluding_day_stay",
        ],
    )


def _header_text(sheet: Worksheet, _layout: CatalogLayout) -> None:
    sheet["E3"] = "Wartość punktowa - hospitalizacja (zmieniona)"


def _scope_label(sheet: Worksheet, layout: CatalogLayout) -> None:
    _ = sheet.cell(4, 13 if layout.catalog == "1a" else 11, "alergologia")


def _extra_header(sheet: Worksheet, _layout: CatalogLayout) -> None:
    _ = sheet.cell(3, 90, "Nowa kolumna")


def _title(sheet: Worksheet, _layout: CatalogLayout) -> None:
    sheet["A2"] = "Katalog świadczeń"


def _legend(sheet: Worksheet, _layout: CatalogLayout) -> None:
    footer = sheet.max_row
    _ = sheet.cell(footer, 1, "3 - inna definicja znacznika")


def _missing_legend(sheet: Worksheet, _layout: CatalogLayout) -> None:
    # openpyxl's cell(row, column, None) leaves a value unchanged.
    sheet.cell(sheet.max_row, 1).value = None


def _after_legend(sheet: Worksheet, _layout: CatalogLayout) -> None:
    _ = sheet.cell(sheet.max_row + 1, 2, "orphan note")


def _body_merge(sheet: Worksheet, _layout: CatalogLayout) -> None:
    # Two empty scope cells of a group row; group cells stay readable.
    sheet.merge_cells("M7:N7")


def _header_merge(sheet: Worksheet, _layout: CatalogLayout) -> None:
    sheet.unmerge_cells("B3:B4")


def _unknown_row(sheet: Worksheet, _layout: CatalogLayout) -> None:
    _ = sheet.cell(sheet.max_row - len(LEGEND), 2, "orphan note")


def _replace_with_text(sheet: Worksheet, layout: CatalogLayout, row: int) -> None:
    for column in range(1, column_index(layout.last_column) + 1):
        sheet.cell(row, column).value = None
    _ = sheet.cell(row, 1, "Uwaga: grupy poniżej")


def _unmerged_chapter(sheet: Worksheet, layout: CatalogLayout) -> None:
    _replace_with_text(sheet, layout, 8)


def _empty_chapter(sheet: Worksheet, layout: CatalogLayout) -> None:
    for row in (7, 8):
        _replace_with_text(sheet, layout, row)
        sheet.merge_cells(f"A{row}:{layout.last_column}{row}")


def _group_before_chapter(sheet: Worksheet, layout: CatalogLayout) -> None:
    sheet.unmerge_cells(f"A6:{layout.last_column}6")
    sheet.cell(6, 1).value = None


def _error_cell(sheet: Worksheet, _layout: CatalogLayout) -> None:
    sheet["E7"] = "#N/A"


def _date_cell(sheet: Worksheet, _layout: CatalogLayout) -> None:
    sheet["E7"] = datetime(2026, 1, 1)


@pytest.mark.parametrize(
    ("edit", "message"),
    [
        (_header_text, "header layout is unsupported"),
        (_scope_label, "scope columns are unsupported"),
        (_extra_header, "header extends beyond the reviewed layout"),
        (_title, "title is unsupported"),
        (_legend, "footer legend is unsupported"),
        (_missing_legend, "footer legend is missing"),
        (_after_legend, "footer legend is unsupported"),
        (_body_merge, "body merge is unsupported"),
        (_header_merge, "header merges are unsupported"),
        (_unknown_row, "row 12 is unsupported"),
        (_unmerged_chapter, "chapter heading is not a merged title row"),
        (_empty_chapter, "chapter before row 7 is empty"),
        (_group_before_chapter, "row 7 is unsupported"),
        (_error_cell, "formula, error or date"),
        (_date_cell, "formula, error or date"),
    ],
)
def test_catalog_refuses_any_unreviewed_layout_drift(
    tmp_path: Path, edit: Callable[[Worksheet, CatalogLayout], None], message: str
) -> None:
    content = catalog_workbook(default_sheets(), edit=edit)
    with pytest.raises(SourceError, match=message):
        _ = parse(tmp_path, content)


def test_shifted_columns_fail_instead_of_misreading_values(tmp_path: Path) -> None:
    # An inserted column moves every value right; reading by position would
    # silently turn hospitalization points into planned-hospitalization points.
    content = catalog_workbook(default_sheets(), shift_from="E")
    # Either header guard may fire first: moved header text or a header cell
    # pushed past the reviewed last column.
    with pytest.raises(SourceError, match="catalog header"):
        _ = parse(tmp_path, content)


def test_original_layout_rejects_a_consolidated_header(tmp_path: Path) -> None:
    content = catalog_workbook({"nfz-1a-46-2026": ROWS_1A})
    manifest = stage_catalog(tmp_path, content, sheets={"nfz-1a-46-2026": ROWS_1A})
    assert len(parse_catalog_workbook(BytesIO(content), manifest)[0].groups) == 5
    wrong = manifest.model_copy(
        update={
            "sheets": {
                "1a": manifest.sheets["1a"].model_copy(
                    update={"layout": "nfz-1a-74-2026"}
                )
            }
        }
    )
    with pytest.raises(SourceError, match="header layout is unsupported"):
        _ = parse_catalog_workbook(BytesIO(content), wrong)


@pytest.mark.parametrize(
    ("values", "message"),
    [
        ({"E": 0}, "not positive"),
        ({"E": -10}, "not positive"),
        ({"E": 100.125}, "two decimals"),
        ({"E": -5.5}, "not positive"),
        ({"E": 10**20}, "exceeds its bound"),
        ({"E": 1e300}, "exceeds its bound"),
        ({"E": 100, "H": 400}, "exceeds its bound"),
        ({"E": "brak"}, "value is unsupported"),
        ({"E": 100, "H": 2.5}, "value is unsupported"),
        ({"E": True}, "value is unsupported"),
        ({"E": "=1+2"}, "formula"),
        ({"H": 7}, "no base points"),
    ],
)
def test_catalog_values_must_be_exact_positive_points(
    tmp_path: Path, values: dict[str, object], message: str
) -> None:
    rows = (Row(1, "A01", "5.51.01.0001001", "Zabiegi", values),)
    content = catalog_workbook({"nfz-1a-74-2026": rows})
    manifest = stage_catalog(
        tmp_path, content, sheets={"nfz-1a-74-2026": rows}, repeated={"1a": 0}
    )
    with pytest.raises(SourceError, match=message):
        _ = parse_catalog_workbook(BytesIO(content), manifest)


@pytest.mark.parametrize(
    ("row", "message"),
    [
        (Row(1, "1A01", "5.51.01.0001001", "x", {"E": 1}), "group code"),
        (Row(1, "A01", "5.51.02.0001001", "x", {"E": 1}), "another catalog"),
        (Row(1, "A01", "5.51.01.001", "x", {"E": 1}), "product code"),
        (Row(1, "A01", "5.51.01.0001001", " ", {"E": 1}), "name is missing"),
        (Row(0, "A01", "5.51.01.0001001", "x", {"E": 1}), "ordinal"),
        (Row(1, "A01", "5.51.01.0001001", "x", {"E": 1}, {3: 4}), "scope marker"),
        (Row(1, "A01", "5.51.01.0001001", "x", {"E": 1}, notes=5), "note"),
    ],
)
def test_catalog_identity_cells_are_validated(
    tmp_path: Path, row: Row, message: str
) -> None:
    content = catalog_workbook({"nfz-1a-74-2026": (row,)})
    manifest = stage_catalog(
        tmp_path, content, sheets={"nfz-1a-74-2026": (row,)}, repeated={"1a": 0}
    )
    with pytest.raises(SourceError, match=message):
        _ = parse_catalog_workbook(BytesIO(content), manifest)


def test_reviewed_counts_and_code_correspondence_fail_closed(tmp_path: Path) -> None:
    content = catalog_workbook(default_sheets())
    with pytest.raises(SourceError, match="repeated catalog codes"):
        _ = parse(tmp_path, content, repeated={"1a": 0})
    manifest = stage_catalog(tmp_path, content)
    fewer = manifest.model_copy(
        update={
            "sheets": {
                **manifest.sheets,
                "1ae": manifest.sheets["1ae"].model_copy(update={"expected_groups": 1}),
            }
        }
    )
    with pytest.raises(SourceError, match="group count"):
        _ = parse_catalog_workbook(BytesIO(content), fewer)
    for rows, message in (
        (ROWS_1A + (ROWS_1A[2],), "more than twice"),
        (
            ROWS_1A[:4] + (Row(5, "A01", "5.51.01.0009999", "x", {"E": 1}),),
            "correspond",
        ),
    ):
        sheets: dict[str, tuple[Row, ...]] = {"nfz-1a-74-2026": rows}
        content = catalog_workbook(sheets)
        manifest = stage_catalog(tmp_path, content, sheets=sheets, repeated={"1a": 2})
        with pytest.raises(SourceError, match=message):
            _ = parse_catalog_workbook(BytesIO(content), manifest)
    shared: dict[str, tuple[Row, ...]] = {
        "nfz-1a-74-2026": ROWS_1A,
        "nfz-1ae-74-2026": (Row(1, "A01", "5.51.02.0001001", "x", {"E": 1}),),
    }
    content = catalog_workbook(shared)
    manifest = stage_catalog(tmp_path, content, sheets=shared)
    with pytest.raises(SourceError, match="two catalogs"):
        _ = parse_catalog_workbook(BytesIO(content), manifest)


def test_workbook_preview_thumbnail_is_allowed_only_by_exact_name(
    tmp_path: Path,
) -> None:
    with ZipFile(BytesIO(catalog_workbook(default_sheets()))) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    with_thumbnail = package_bytes(
        {**parts, "docProps/thumbnail.emf": b"\x01\x00\x00\x00"}
    )
    assert len(parse(tmp_path, with_thumbnail)) == len(ROWS_1A) + len(ROWS_1AE)
    other_image = package_bytes({**parts, "docProps/preview.emf": b"\x01\x00\x00\x00"})
    with pytest.raises(SourceError, match="part type"):
        _ = parse(tmp_path, other_image)


def test_manifest_binds_layouts_catalogs_and_lineage(tmp_path: Path) -> None:
    manifest = stage_catalog(tmp_path, catalog_workbook(default_sheets()))
    data = manifest.model_dump(mode="python")
    for change, message in (
        ({"sheets": {"1a": {**data["sheets"]["1a"], "layout": "nfz-1ae-74-2026"}}}, "not reviewed"),
        ({"sheets": {"1ae": data["sheets"]["1ae"]}}, "must contain catalog 1a"),
        ({"introductory_document_id": None}, "requires its lineage"),
        ({"automatic_type": 1}, "original act has no source"),
        ({"document_url": "https://baw.nfz.gov.pl/NFZ/document/1/Other"}, "does not match"),
        ({"workbook": {**data["workbook"], "source_url": "https://example.com/1a.xlsx"}}, "source_url"),
    ):  # fmt: skip
        with pytest.raises(ValidationError, match=message):
            _ = CatalogManifest.model_validate({**data, **change}, strict=False)


def test_blank_notes_are_absent_not_text(tmp_path: Path) -> None:
    rows = (Row(1, "A01", "5.51.01.0001001", "x", {"E": 1}, notes=" \xa0 "),)
    content = catalog_workbook({"nfz-1a-74-2026": rows})
    manifest = stage_catalog(tmp_path, content, sheets={"nfz-1a-74-2026": rows})
    sheet = parse_catalog_workbook(BytesIO(content), manifest)[0]
    assert sheet.groups[0].notes is None
    assert sheet.chapters == ("Choroby układu nerwowego",)


def test_layouts_must_assign_every_column_exactly_once() -> None:
    gap = replace(NFZ_1A_74_2026, measures=NFZ_1A_74_2026.measures[:-1])
    with pytest.raises(ValueError, match="covered exactly once"):
        _ = check_layout(gap)
    narrow = replace(NFZ_1A_74_2026, last_scope_column="CA", notes_column="CB")
    with pytest.raises(ValueError, match="scope columns do not match"):
        _ = check_layout(narrow)
