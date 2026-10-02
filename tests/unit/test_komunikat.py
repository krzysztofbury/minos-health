from datetime import date
from io import BytesIO
from pathlib import Path
from typing import cast
from zipfile import ZipFile

import pytest
from pydantic import ValidationError

from minos.domain.komunikat import DepartmentSpecialty, ProcedureRank
from minos.domain.reference import SourceError
from minos.ingest.komunikat import parse_algorithm, parse_companion
from minos.ingest.komunikat_manifest import KomunikatManifest
from minos.ingest.workbook_archive import WorkbookLimits
from tests.fixtures.komunikat import (
    PARAGRAPHS,
    RANK_ROWS,
    companion_bytes,
    docx_bytes,
    package_bytes,
    stage_komunikat,
)


@pytest.fixture
def manifest(tmp_path: Path) -> KomunikatManifest:
    return stage_komunikat(tmp_path, "icd9-czs-synthetic-v1")


def test_algorithm_extracts_split_runs_and_reporting_month_not_publication(
    manifest: KomunikatManifest,
) -> None:
    metadata = parse_algorithm(BytesIO(docx_bytes()), manifest)
    assert metadata.version == "1.1.2026"
    assert metadata.reporting_from == date(2026, 1, 1)
    assert metadata.version_text == "wersja 1.1.2026"
    assert metadata.paragraph == 2


def test_algorithm_reporting_month_is_y_not_version_number_x(
    manifest: KomunikatManifest,
) -> None:
    # 1.1.2026 cannot tell X from Y; a distinct pair pins the source semantics.
    paragraphs = PARAGRAPHS[:1] + ("wersja 2.4.2026",) + PARAGRAPHS[2:]
    metadata = parse_algorithm(
        BytesIO(docx_bytes(paragraphs)),
        manifest.model_copy(update={"algorithm_version": "2.4.2026"}),
    )
    assert metadata.reporting_from == date(2026, 4, 1)


@pytest.mark.parametrize(
    "paragraphs",
    [
        PARAGRAPHS + ("wersja 2.1.2026",),
        PARAGRAPHS[:1] + ("wersja 1.13.2026",) + PARAGRAPHS[2:],
        PARAGRAPHS[:1] + ("wersja 2.1.2026",) + PARAGRAPHS[2:],
        PARAGRAPHS[:3],
        PARAGRAPHS[:3] + (PARAGRAPHS[3] + " nie obowiązuje",) + PARAGRAPHS[4:],
        PARAGRAPHS[:4] + PARAGRAPHS[5:],
        PARAGRAPHS[:-1],
    ],
)
def test_algorithm_rejects_ambiguous_version_invalid_month_drift_and_missing_semantics(
    manifest: KomunikatManifest, paragraphs: tuple[str, ...]
) -> None:
    with pytest.raises(SourceError, match="algorithm"):
        _ = parse_algorithm(BytesIO(docx_bytes(paragraphs)), manifest)


def test_companion_preserves_code_types_titles_and_scope_without_filling_blank_rows(
    manifest: KomunikatManifest,
) -> None:
    rows = list(
        parse_companion(BytesIO(companion_bytes(ranks=True)), manifest, ranks=True)
    )
    assert rows == [
        ProcedureRank("SZP", "A01", "  Synthetic rank title  ", 4, "SZP 2026", 4),
        ProcedureRank("SZP", "00.00001", "Recorded zero", 0, "SZP 2026", 5),
        ProcedureRank("SZP", "05.01", "Not a canonical concept", 3, "SZP 2026", 6),
        ProcedureRank("AOS", "A01", "Other grouper rank", 2, "AOS 2026", 4),
    ]
    specialties = list(
        parse_companion(BytesIO(companion_bytes(ranks=False)), manifest, ranks=False)
    )
    assert specialties[0] == DepartmentSpecialty(
        "SZP", "0010", "4", "WSZYSTKIE", "SZP 2026", 4
    )
    assert len(specialties) == 3


@pytest.mark.parametrize(
    "row",
    [
        ("A01", "Blank", True),
        ("A01", "Float", 3.5),
        ("A01", "Unknown", 5),
        ("A01", "Formula", "=1+2"),
        (1.1, "Ambiguous numeric code", 3),
        ("A01", None, 3),
    ],
)
def test_rank_cells_reject_unsupported_types_and_values(
    manifest: KomunikatManifest, row: tuple[object, ...]
) -> None:
    content = companion_bytes(ranks=True, rows={**RANK_ROWS, "SZP": (row,)})
    with pytest.raises(SourceError, match="row value|formula"):
        _ = list(parse_companion(BytesIO(content), manifest, ranks=True))


def test_rank_duplicate_and_reviewed_count_drift_fail(
    manifest: KomunikatManifest,
) -> None:
    for rows, expected in (
        ({**RANK_ROWS, "SZP": RANK_ROWS["SZP"] + (RANK_ROWS["SZP"][0],)}, "duplicated"),
        ({**RANK_ROWS, "SZP": RANK_ROWS["SZP"][:1]}, "row count"),
    ):
        with pytest.raises(SourceError, match=expected):
            _ = list(
                parse_companion(
                    BytesIO(companion_bytes(ranks=True, rows=rows)),
                    manifest,
                    ranks=True,
                )
            )


@pytest.mark.parametrize(
    "drift", ["title", "header", "scope", "merge", "extra_column", "partial_row"]
)
def test_companion_refuses_unreviewed_table_layout(
    manifest: KomunikatManifest, drift: str
) -> None:
    from openpyxl import load_workbook
    from openpyxl.worksheet.worksheet import Worksheet

    book = load_workbook(BytesIO(companion_bytes(ranks=True)))
    sheet = cast(Worksheet, book["SZP 2026"])
    if drift == "title":
        sheet["A1"] = "Changed table title"
    elif drift == "header":
        sheet["C2"] = "Other meaning"
    elif drift == "scope":
        sheet.title = "REH 2026"
    elif drift == "merge":
        sheet.merge_cells("A1:C1")
    elif drift == "extra_column":
        sheet["D4"] = "Unexpected additional field"
    else:
        sheet.append((None, "Not an empty styled row", 3))
    output = BytesIO()
    book.save(output)
    book.close()
    with pytest.raises(SourceError, match="layout|identities|row value"):
        _ = list(parse_companion(BytesIO(output.getvalue()), manifest, ranks=True))


def test_companion_maps_reader_structure_failure_to_source_error(
    manifest: KomunikatManifest,
) -> None:
    with ZipFile(BytesIO(companion_bytes(ranks=True))) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["xl/styles.xml"] = (
        b'<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><fonts><font><sz val="not-a-number"/></font></fonts></styleSheet>'
    )
    with pytest.raises(SourceError, match="reader rejected source structure"):
        _ = list(parse_companion(BytesIO(package_bytes(parts)), manifest, ranks=True))


@pytest.mark.parametrize(
    "alteration", ["dtd", "external", "duplicate", "traversal", "extra_depth", "ratio"]
)
def test_docx_uses_zip_and_xml_security_guards(
    manifest: KomunikatManifest, alteration: str
) -> None:
    with ZipFile(BytesIO(docx_bytes())) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    if alteration == "dtd":
        parts["word/document.xml"] = (
            b'<!DOCTYPE x [<!ENTITY test "unsafe">]><x>&test;</x>'
        )
    elif alteration == "external":
        parts["_rels/.rels"] = (
            b'<Relationships><Relationship TargetMode="External" Target="https://invalid.example"/></Relationships>'
        )
    elif alteration == "duplicate":
        parts["word/document.xml"] = (
            b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body/><w:body/></w:document>'
        )
    elif alteration == "traversal":
        parts["../word/unsafe.xml"] = b"<x/>"
    elif alteration == "extra_depth":
        parts["word/document.xml"] = b"<x>" * 40 + b"</x>" * 40
    else:
        parts["word/document.xml"] = b"<x>" + b"a" * 200_000 + b"</x>"
    with pytest.raises(SourceError):
        _ = parse_algorithm(BytesIO(package_bytes(parts)), manifest)


def test_docx_entry_and_source_limits_apply_before_extraction(
    manifest: KomunikatManifest,
) -> None:
    for limits in (WorkbookLimits(max_entries=2), WorkbookLimits(max_source_bytes=64)):
        with pytest.raises(SourceError, match="limit"):
            _ = parse_algorithm(BytesIO(docx_bytes()), manifest, limits=limits)


def test_komunikat_manifest_refuses_mixed_notice_and_unreviewed_pairing(
    manifest: KomunikatManifest,
) -> None:
    data = manifest.model_dump(mode="python")
    data["pairing"] = "ruleset"
    with pytest.raises(ValidationError, match="pairing"):
        _ = KomunikatManifest.model_validate(data)
    data = manifest.model_dump(mode="python")
    data["ranks"] = manifest.ranks.model_copy(
        update={"source_url": manifest.ranks.source_url.replace("/54/", "/55/")}
    )
    with pytest.raises(ValidationError, match="different notices"):
        _ = KomunikatManifest.model_validate(data)
